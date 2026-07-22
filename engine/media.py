"""Media conversion: PDF / DOCX / video → the CV pipeline's native inputs.

The CV-NLP layers analyse exactly three sources (image list, text, html).
Rather than teach every layer new formats, richer media converts at the
engine boundary:

    pdf   -> one rendered PNG per page   -> the multi-image (carousel) path,
             so OCR geometry, contrast and disclosure-proximity measurements
             all keep working per page
    docx  -> extracted paragraph + table text -> the text path
    video -> uniformly sampled keyframes -> the multi-image path; the
             card->timestamp map rides in provenance.media so findings can be
             traced back to a moment in the ad

Page/frame counts are bounded (upload validation enforces its own caps
first) so a hostile or huge file cannot amplify OCR cost.
"""

from __future__ import annotations

import os

PDF_RENDER_DPI = 150          # OCR-friendly render without huge bitmaps
PDF_MAX_PAGES = 40            # mirrors the API's upload cap
VIDEO_MAX_FRAMES = 12         # per run; ~one frame per 5 s up to this many
VIDEO_FRAME_EVERY_S = 5.0


def pdf_to_images(pdf_path: str, workdir: str) -> tuple[list[str], dict]:
    """Render each PDF page to a PNG in workdir. → (paths, media meta)."""
    import fitz  # PyMuPDF

    paths: list[str] = []
    with fitz.open(pdf_path) as doc:
        if doc.needs_pass:
            raise ValueError("password-protected PDFs cannot be analysed")
        total_pages = doc.page_count
        count = min(total_pages, PDF_MAX_PAGES)
        for i in range(count):
            pix = doc.load_page(i).get_pixmap(dpi=PDF_RENDER_DPI)
            p = os.path.join(workdir, f"page_{i + 1:03d}.png")
            pix.save(p)
            paths.append(p)
    if not paths:
        raise ValueError("the PDF has no renderable pages")
    return paths, {"kind": "pdf", "pages": len(paths),
                   "truncated": total_pages > len(paths)}


def docx_to_text(docx_path: str) -> tuple[str, dict]:
    """Word document → plain text (paragraphs, then tables row-joined)."""
    import docx

    d = docx.Document(docx_path)
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), {"kind": "docx", "blocks": len(parts)}


def video_to_frames(video_path: str, workdir: str,
                    max_frames: int = VIDEO_MAX_FRAMES,
                    every_s: float = VIDEO_FRAME_EVERY_S
                    ) -> tuple[list[str], dict]:
    """Sample keyframes uniformly across the video. → (paths, media meta).

    Frames become carousel cards (card N = frames[N-1]); the meta's
    frames[] list maps each card to its timestamp for report display.
    """
    import cv2

    cap = cv2.VideoCapture(video_path)
    try:
        if not cap.isOpened():
            raise ValueError("the video could not be decoded")
        fps = cap.get(cv2.CAP_PROP_FPS) or 0.0
        total = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0.0
        if fps <= 0 or total <= 0:
            raise ValueError("the video reports no readable duration")
        duration = float(total) / float(fps)
        count = max(1, min(max_frames, int(duration // every_s) + 1))
        paths: list[str] = []
        frames_meta: list[dict] = []
        for i in range(count):
            t = duration * (i + 0.5) / count      # sample mid-slot
            cap.set(cv2.CAP_PROP_POS_MSEC, t * 1000.0)
            ok, frame = cap.read()
            if not ok or frame is None:
                continue
            p = os.path.join(workdir, f"frame_{i + 1:03d}.png")
            if cv2.imwrite(p, frame):
                paths.append(p)
                frames_meta.append({"card": len(paths), "t": round(t, 2)})
        if not paths:
            raise ValueError("no frames could be extracted from the video")
        return paths, {"kind": "video",
                       "duration_seconds": round(duration, 2),
                       "sampled_frames": len(paths),
                       "frames": frames_meta}
    finally:
        cap.release()
